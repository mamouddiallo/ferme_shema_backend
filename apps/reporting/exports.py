"""
Export du rapport mensuel (§10, backlog R2). reportlab pour le PDF (pur
Python, pas de dépendance système contrairement à weasyprint — plus simple
à déployer en conteneur) et openpyxl pour l'Excel.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MOIS_FR = {
    1: "Janvier",
    2: "Février",
    3: "Mars",
    4: "Avril",
    5: "Mai",
    6: "Juin",
    7: "Juillet",
    8: "Août",
    9: "Septembre",
    10: "Octobre",
    11: "Novembre",
    12: "Décembre",
}


def _mois_fr(d) -> str:
    """Nom du mois en français, sans dépendre de la locale système (absente par défaut en conteneur)."""
    return f"{MOIS_FR[d.month]} {d.year}"


def generer_pdf_rapport_mensuel(rapport: dict) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Ferme SHEMA — Rapport mensuel", styles["Title"]))
    elements.append(Paragraph(f"Mois : {_mois_fr(rapport['mois'])}", styles["Normal"]))
    elements.append(Spacer(1, 0.8 * cm))

    donnees_synthese = [
        ["Indicateur", "Montant (FCFA)"],
        ["Chiffre d'affaires", f"{rapport['chiffre_affaires']:,.0f}"],
        ["Total des dépenses", f"{rapport['total_depenses']:,.0f}"],
        ["Marge", f"{rapport['marge']:,.0f}"],
        ["Nombre de ventes", str(rapport["nombre_ventes"])],
    ]
    table_synthese = Table(donnees_synthese, colWidths=[8 * cm, 6 * cm])
    table_synthese.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f6e56")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f6f2")]),
            ]
        )
    )
    elements.append(table_synthese)
    elements.append(Spacer(1, 1 * cm))

    if rapport["depenses_par_categorie"]:
        elements.append(Paragraph("Détail des dépenses par catégorie", styles["Heading2"]))
        donnees_depenses = [["Catégorie", "Montant (FCFA)"]]
        for categorie, montant in rapport["depenses_par_categorie"].items():
            donnees_depenses.append([categorie, f"{montant:,.0f}"])
        table_depenses = Table(donnees_depenses, colWidths=[8 * cm, 6 * cm])
        table_depenses.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3c3489")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ]
            )
        )
        elements.append(table_depenses)

    doc.build(elements)
    return buffer.getvalue()


def generer_excel_rapport_mensuel(rapport: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport mensuel"

    gras = Font(bold=True)

    ws["A1"] = "Ferme SHEMA — Rapport mensuel"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Mois : {_mois_fr(rapport['mois'])}"

    ws["A4"] = "Indicateur"
    ws["B4"] = "Montant (FCFA)"
    ws["A4"].font = gras
    ws["B4"].font = gras

    lignes = [
        ("Chiffre d'affaires", float(rapport["chiffre_affaires"])),
        ("Total des dépenses", float(rapport["total_depenses"])),
        ("Marge", float(rapport["marge"])),
        ("Nombre de ventes", rapport["nombre_ventes"]),
    ]
    for i, (libelle, valeur) in enumerate(lignes, start=5):
        ws[f"A{i}"] = libelle
        ws[f"B{i}"] = valeur

    ligne_suivante = 5 + len(lignes) + 1
    if rapport["depenses_par_categorie"]:
        ws[f"A{ligne_suivante}"] = "Dépenses par catégorie"
        ws[f"A{ligne_suivante}"].font = gras
        ligne_suivante += 1
        for categorie, montant in rapport["depenses_par_categorie"].items():
            ws[f"A{ligne_suivante}"] = categorie
            ws[f"B{ligne_suivante}"] = float(montant)
            ligne_suivante += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
